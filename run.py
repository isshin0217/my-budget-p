"""
run.py
카드 문자를 파싱해 엑셀 가계부에 추가하고, 이번 주 현황을 바로 보여준다.

사용법:
  python run.py messages.txt   문자들을 붙여넣은 파일에서 읽기
  type messages.txt | python run.py
  python run.py --demo         데모 데이터로 만들어 보기
  python run.py --week         새 문자 없이 이번 주 현황만 다시 보기

결과: 가계부.xlsx (내역 / 주간요약 / 월간요약 / 설정)
"""

import re
import sys
from pathlib import Path

# 파이프로 넘길 때 한글이 깨지지 않도록
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook

from parser import parse_message
from ledger import SHEET, append_rows, ensure_settings, migrate, read_settings
from report import build_all, load_records

XLSX = "가계부.xlsx"

DEMO = """[Web 발신]
신한카드(1234)승인 홍*동 94,400원(일시불)07/01 19:55 폴리스 유한

[Web 발신]
신한카드(1234)승인 홍*동 5,500원(일시불)07/02 08:12 스타벅스 강남

[Web 발신]
신한카드(1234)취소 홍*동 12,000원(일시불)07/03 21:40 배달의민족
"""


def split_messages(blob: str) -> list[str]:
    """빈 줄과 '[Web 발신]' 머리말을 경계로 문자 블록을 나눈다.

    빈 줄만 보면 안 된다. 붙여넣다가 줄 하나가 빠지는 순간 두 건이 한 덩어리가
    되는데, 정규식은 앞에서부터 훑다가 뒤엣것 하나만 잡거나 뒤 문자를 가맹점명
    안으로 통째로 삼킨다. 어느 쪽이든 앞의 결제가 조용히 사라지고, 실패 목록에도
    안 뜬다 — '파싱 성공 1건' 이라고 찍히니 알아챌 방법이 없다.

    카드 문자에는 거의 예외 없이 이 머리말이 붙으므로 여기를 경계로 삼는다.
    (머리말이 없는 문자 둘이 빈 줄 없이 붙는 경우까지는 못 가른다.)
    """
    blob = re.sub(r"\[Web\s*발신\]", "\n\n[Web 발신]", blob)
    blocks, cur = [], []
    for line in blob.splitlines():
        if line.strip() == "":
            if cur:
                blocks.append("\n".join(cur)); cur = []
        else:
            cur.append(line)
    if cur:
        blocks.append("\n".join(cur))
    return blocks


def _warn_dropped(dropped: list):
    """날짜를 못 읽어 리포트에서 빠진 행을 알린다.

    시트에는 남아 있는데 합계에만 안 잡히므로, 말해주지 않으면 왜 숫자가
    안 맞는지 알아낼 방법이 없다.
    """
    if not dropped:
        return
    print(f"날짜를 못 읽어 리포트에서 뺀 행 {len(dropped)}건 "
          f"(시트에는 그대로 있습니다. 날짜시간 칸을 'YYYY-MM-DD HH:MM' 으로 고쳐주세요):")
    for v in dropped[:5]:
        print("   ·", str(v)[:40])


def print_week(s: dict):
    """이번 주 현황. 선결제 전에 이 두 줄만 보면 된다."""
    bar_len = 30
    used = s["variable"]
    allow = s["weekly_allow"] or 1
    filled = max(0, min(bar_len, round(bar_len * used / allow)))
    bar = "■" * filled + "·" * (bar_len - filled)   # cp949 콘솔에서도 안전한 문자
    over = s["remaining"] < 0

    print()
    print(f"  이번 주  {s['start']:%m/%d} ~ {s['end']:%m/%d}   (남은 {s['days_left']}일)")
    print(f"  {bar}  {used:,} / {s['weekly_allow']:,}원")
    print()
    if over:
        print(f"  이번 주 {-s['remaining']:,}원 초과했습니다.")
    else:
        print(f"  더 써도 되는 돈   {s['remaining']:,}원   (하루 {s['per_day']:,}원 페이스)")
    print(f"  지금 선결제할 금액 {s['charged']:,}원")
    print()
    if s["fixed_total"]:
        names = ", ".join(list(s["detected"])[:3]) or "설정값"
        print(f"  · 월 고정비 {s['fixed_total']:,}원 ({names}) 은 주간 한도에서 미리 뺐습니다.")
    if not s["budget"]:
        print("  · '설정' 시트에 월 예산을 넣으면 남은 돈이 계산됩니다.")


def main():
    args = [a for a in sys.argv[1:]]

    if "--week" in args:
        if not Path(XLSX).exists():
            print(f"{XLSX} 이 없습니다. 먼저 문자를 넣어주세요.")
            return
        wb = load_workbook(XLSX)
        if SHEET not in wb.sheetnames:
            print(f"'{SHEET}' 시트가 없습니다. 먼저 문자를 넣어주세요.")
            return
        ensure_settings(wb)
        migrate(wb[SHEET])
        dropped = []
        summary = build_all(wb, load_records(wb, dropped), read_settings(wb))
        wb.save(XLSX)
        _warn_dropped(dropped)
        print_week(summary)
        return

    if "--demo" in args:
        blob = DEMO
    elif args and not args[0].startswith("--"):
        blob = open(args[0], encoding="utf-8").read()
    else:
        blob = sys.stdin.read()

    records, failed = [], []
    for block in split_messages(blob):
        # 한 건이 터졌다고 나머지까지 잃으면 안 된다. 실패 목록으로 보내고 계속한다.
        try:
            r = parse_message(block)
        except Exception as err:
            print(f"   ! 읽다가 오류: {err}")
            r = None
        (records if r else failed).append(r if r else block)

    added, wb, path = append_rows(records, XLSX)
    dropped = []
    summary = build_all(wb, load_records(wb, dropped), read_settings(wb))
    wb.save(path)

    print(f"파싱 성공 {len(records)}건 · 신규 저장 {added}건 → {XLSX}")
    if failed:
        print(f"파싱 실패 {len(failed)}건 (형식이 다를 수 있음):")
        for f in failed:
            print("   ·", f.replace(chr(10), " ")[:70])
    _warn_dropped(dropped)
    print_week(summary)


if __name__ == "__main__":
    main()
