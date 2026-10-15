"""
report.py
내역 시트를 읽어 요약 리포트를 만든다. 중심은 '이번 주에 얼마 더 써도 되는가'.

계산 모델
  월 고정비   : 같은 가맹점이 약 한 달 간격으로 비슷한 금액 반복 → 자동 감지
  주간 가용액 = (월 예산 − 월 고정비) ÷ 4.35
  이번 주 남은 돈 = 주간 가용액 − 이번 주 변동비
  선결제 금액 = 이번 주 카드 승인 총액 (고정비 포함)

고정비는 월 예산에서 이미 뺐으므로 주간 한도에서 또 빼지 않는다.
다만 카드로는 실제로 나갔으므로 선결제 금액에는 포함된다.
"""

from collections import defaultdict
from datetime import datetime, timedelta

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from ledger import HEADERS, SHEET

WEEKDAYS = "월화수목금토일"
# 설정 시트에는 한글로도 영어로도 적을 수 있다. Code.gs 쪽 Settings 는 영어라
# 같은 파일을 두 곳에서 쓰면 여기서도 영어를 알아들어야 주 경계가 갈리지 않는다.
WEEKDAYS_EN = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
WEEKS_PER_MONTH = 4.35          # 365 ÷ 7 ÷ 12
# 마지막 결제가 이만큼 지났으면 이미 끊은 것으로 본다 (Code.gs 와 같은 값).
FIXED_FRESH_DAYS = 40


def _money(v) -> int:
    """설정 칸의 금액. '2,000,000' 처럼 콤마를 넣어도 받는다. 못 읽으면 0.

    예전에는 int() 를 그냥 불러서, 사람이 자연스럽게 콤마를 찍어두면 ValueError 로
    실행이 통째로 죽었다. 그러면 그 앞에서 만든 원장도 저장되지 못하고 날아간다.
    """
    try:
        return int(float(str(v if v is not None else 0).replace(",", "").strip() or 0))
    except (TypeError, ValueError):
        return 0


def _parse_dt(v):
    """셀 값 → datetime. 못 읽으면 None.

    엑셀이 날짜값으로 바꿔둔 셀도, 사람이 날짜만 적은 셀도 받는다. 예전에는
    '%Y-%m-%d %H:%M' 하나만 받아서, 손으로 적은 '2026-07-05' 행이 리포트에서
    통째로 빠졌다 — 시트에는 남아 있는데 합계에만 안 잡히니 알아채기 어렵다.
    """
    if isinstance(v, datetime):
        return v
    s = str(v or "").strip()
    for fmt, n in (("%Y-%m-%d %H:%M", 16), ("%Y-%m-%d", 10)):
        try:
            return datetime.strptime(s[:n], fmt)
        except ValueError:
            continue
    return None

MONTHLY = "월간요약"
WEEKLY = "주간요약"

HEAD_FILL = PatternFill("solid", fgColor="0A2342")
HEAD_FONT = Font(color="FFFFFF", bold=True)
BIG_FILL = PatternFill("solid", fgColor="E8F0FE")
OK_FONT = Font(bold=True, size=14, color="1B7F3B")
OVER_FONT = Font(bold=True, size=14, color="C0362C")
WON = '#,##0"원"'


# ---------------------------------------------------------------- 데이터 읽기

def load_records(wb, dropped: list | None = None) -> list[dict]:
    """'내역' 시트를 파이썬 레코드로 읽는다.

    날짜를 못 읽은 행은 버리되, dropped 리스트를 주면 거기에 담아 돌려준다 —
    조용히 사라지면 합계가 왜 안 맞는지 알 길이 없다.
    """
    if SHEET not in wb.sheetnames:
        return []
    idx = {h: i for i, h in enumerate(HEADERS)}
    out = []
    for row in wb[SHEET].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        dt = _parse_dt(row[0])
        if dt is None:
            if dropped is not None:
                dropped.append(row[0])
            continue
        out.append({
            "dt": dt,
            "amount": row[idx["금액"]] or 0,
            "category": row[idx["카테고리"]] or "미분류",
            "merchant": str(row[idx["가맹점"]] or ""),
            "period": row[idx["주기"]] or "주기",
            "tag": str(row[idx["태그"]] or "").strip(),
        })
    return out


# ------------------------------------------------------------- 고정비 자동감지

def detect_fixed(records: list[dict]) -> dict:
    """가맹점별로 한 달 간격 반복 결제를 찾아 {가맹점: 최근금액} 로 돌려준다.

    판정: 연속 결제 간격이 26~35일이고 금액 차이가 15% 이내인 쌍이 하나라도 있으면 고정비.
    데이터가 2개월 이상 쌓여야 잡히므로 그전에는 설정의 '월 고정비(수동)'를 쓴다.
    마지막 결제가 40일을 넘었으면 이미 끊은 것으로 보고 뺀다 — 안 그러면 작년에
    두 번 가고 만 미용실이 지금까지 매달 예산을 깎는다.
    """
    by_merchant = defaultdict(list)
    for r in records:
        if r["amount"] > 0:                       # 취소(음수)는 주기 판정에서 제외
            by_merchant[r["merchant"]].append(r)

    now = datetime.now()
    fixed = {}
    for merchant, rows in by_merchant.items():
        rows.sort(key=lambda r: r["dt"])
        if (now.date() - rows[-1]["dt"].date()).days > FIXED_FRESH_DAYS:
            continue
        for prev, cur in zip(rows, rows[1:]):
            # 달력 날짜로 센다. timedelta.days 는 시각까지 보고 내림하므로
            # 23시 결제 다음이 01시면 26일이 25일로 읽혀 고정비를 놓친다.
            gap = (cur["dt"].date() - prev["dt"].date()).days
            base = max(prev["amount"], cur["amount"]) or 1
            if 26 <= gap <= 35 and abs(cur["amount"] - prev["amount"]) / base <= 0.15:
                fixed[merchant] = rows[-1]["amount"]
                break
    return fixed


def monthly_fixed_total(records, settings) -> tuple[int, dict]:
    """월 고정비 합계와 감지 내역. 설정의 수동값이 0보다 크면 그 값을 우선한다."""
    detected = detect_fixed(records)
    manual = _money(settings.get("월 고정비(수동)"))
    return (manual if manual > 0 else sum(detected.values())), detected


# ------------------------------------------------------------------- 주 계산

def week_start_index(settings) -> int:
    raw = str(settings.get("주 시작 요일", "월")).strip()
    for i, en in enumerate(WEEKDAYS_EN):        # 'Mon' / 'sun' 처럼 영어로 적어도 받는다
        if raw[:3].lower() == en.lower():
            return i
    ch = raw[:1]
    return WEEKDAYS.index(ch) if ch in WEEKDAYS else 0


def week_start_of(d: datetime, start_i: int):
    d = d.date() if isinstance(d, datetime) else d
    return d - timedelta(days=(d.weekday() - start_i) % 7)


def summarize_week(records, settings, today=None) -> dict:
    """이번 주 현황. run.py가 터미널에 찍는 숫자와 주간요약 시트가 같은 값을 쓴다."""
    today = today or datetime.now()
    start_i = week_start_index(settings)
    ws_start = week_start_of(today, start_i)
    ws_end = ws_start + timedelta(days=6)

    budget = _money(settings.get("월 예산"))
    fixed_total, detected = monthly_fixed_total(records, settings)
    weekly_allow = int((budget - fixed_total) / WEEKS_PER_MONTH) if budget else 0

    this_week = [r for r in records if ws_start <= r["dt"].date() <= ws_end]
    charged = sum(r["amount"] for r in this_week)                     # 선결제 대상
    # 변동비에서 빼는 건 '월 예산에서 이미 뺀 것'(= 고정비로 감지된 것)뿐이다.
    # 카테고리가 구독/통신이라는 이유만으로 빼면, 아직 고정비로 안 잡힌 첫 결제가
    # 계산 어디에도 안 남는다.
    variable = sum(r["amount"] for r in this_week if r["merchant"] not in detected)

    remaining = weekly_allow - variable
    days_left = (ws_end - today.date()).days + 1
    days_left = max(days_left, 0)

    return {
        "start": ws_start, "end": ws_end,
        "budget": budget, "fixed_total": fixed_total, "detected": detected,
        "weekly_allow": weekly_allow,
        "charged": charged, "variable": variable,
        "remaining": remaining,
        "days_left": days_left,
        "per_day": int(remaining / days_left) if days_left else 0,
        "rows": this_week,
    }


# ------------------------------------------------------------------- 시트 작성

def _reset_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def _title(ws, row, text):
    c = ws.cell(row, 1, text)
    c.font = Font(bold=True, size=11, color="0A2342")
    return row + 1


def _kv(ws, row, label, value, fmt=WON, big=False):
    ws.cell(row, 1, label)
    c = ws.cell(row, 2, value)
    c.number_format = fmt
    if big:
        ws.cell(row, 1).font = Font(bold=True, size=12)
        c.font = OK_FONT if (isinstance(value, (int, float)) and value >= 0) else OVER_FONT
        ws.cell(row, 1).fill = BIG_FILL
        c.fill = BIG_FILL
    return row + 1


def _breakdown(ws, row, header, pairs, sort_key=None):
    """(이름, 금액) 목록을 쓴다. 기본은 큰 금액순, sort_key 를 주면 그 순서로."""
    row = _title(ws, row, header)
    for name, amount in sorted(pairs, key=sort_key or (lambda x: -x[1])):
        ws.cell(row, 1, name)
        ws.cell(row, 2, amount).number_format = WON
        row += 1
    return row + 1


def build_weekly(wb, records, settings):
    """주간요약: 맨 위에 '이번 주 남은 돈', 아래에 주차별 추이."""
    s = summarize_week(records, settings)
    ws = _reset_sheet(wb, WEEKLY)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    for col in "CDEFG":
        ws.column_dimensions[col].width = 14

    row = _title(ws, 1, f"이번 주 ({s['start']:%m/%d} ~ {s['end']:%m/%d})")
    row = _kv(ws, row, "이번 주 남은 돈", s["remaining"], big=True)
    row = _kv(ws, row, "지금 선결제할 금액", s["charged"], big=True)
    row += 1
    row = _kv(ws, row, "주간 가용액", s["weekly_allow"])
    row = _kv(ws, row, "이번 주 변동비", s["variable"])
    row = _kv(ws, row, "남은 일수", s["days_left"], fmt='0"일"')
    row = _kv(ws, row, "하루 페이스", s["per_day"])
    row += 1
    row = _kv(ws, row, "월 예산", s["budget"])
    row = _kv(ws, row, "월 고정비", s["fixed_total"])
    row += 1

    if s["detected"]:
        row = _breakdown(ws, row, "자동 감지된 고정비", list(s["detected"].items()))

    by_tag = defaultdict(int)
    for r in s["rows"]:
        by_tag[r["tag"] or "(태그 없음)"] += r["amount"]
    if by_tag:
        row = _breakdown(ws, row, "이번 주 태그별", list(by_tag.items()))

    # 주차별 추이
    start_i = week_start_index(settings)
    weeks = defaultdict(lambda: {"총액": 0, "변동비": 0})
    for r in records:
        k = week_start_of(r["dt"], start_i)
        weeks[k]["총액"] += r["amount"]
        # 위의 '이번 주 변동비'와 **같은 기준**으로 센다. 예전에는 여기에만
        # period != '고정' 을 더 걸어서, 같은 주가 시트 위아래에서 서로 다른
        # 숫자로 나왔다 — 아직 고정비로 안 잡힌 통신비 한 건에 5만원씩 벌어졌다.
        # 카테고리가 '고정'이라고 빼면 안 되는 이유는 summarize_week 주석에.
        if r["merchant"] not in s["detected"]:
            weeks[k]["변동비"] += r["amount"]

    row = _title(ws, row, "주차별 추이")
    head_row = row
    for i, h in enumerate(["주 시작", "변동비", "주간 한도", "잔여"], 1):
        c = ws.cell(row, i, h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center")
    row += 1
    for k in sorted(weeks):
        ws.cell(row, 1, f"{k:%Y-%m-%d}")
        ws.cell(row, 2, weeks[k]["변동비"]).number_format = WON
        ws.cell(row, 3, s["weekly_allow"]).number_format = WON
        ws.cell(row, 4, s["weekly_allow"] - weeks[k]["변동비"]).number_format = WON
        row += 1

    if len(weeks) >= 2:
        chart = BarChart()
        chart.title = "주차별 변동비 vs 주간 한도"
        chart.height, chart.width = 7, 16
        data = Reference(ws, min_col=2, max_col=3, min_row=head_row, max_row=row - 1)
        cats = Reference(ws, min_col=1, min_row=head_row + 1, max_row=row - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"F{head_row}")

    return s


def build_monthly(wb, records, settings):
    """월간요약: 이번 달 주기별·카테고리별·태그별 합계와 월별 추이."""
    ws = _reset_sheet(wb, MONTHLY)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16

    now = datetime.now()
    this_month = [r for r in records if (r["dt"].year, r["dt"].month) == (now.year, now.month)]
    _, detected = monthly_fixed_total(records, settings)
    budget = _money(settings.get("월 예산"))
    spent = sum(r["amount"] for r in this_month)

    row = _title(ws, 1, f"{now:%Y년 %m월}")
    row = _kv(ws, row, "이번 달 지출", spent, big=True)
    if budget:
        row = _kv(ws, row, "예산 대비 남은 돈", budget - spent, big=True)
    row += 1

    by_period = defaultdict(int)
    by_category = defaultdict(int)
    by_tag = defaultdict(int)
    for r in this_month:
        period = "고정" if r["merchant"] in detected else r["period"]
        by_period[period] += r["amount"]
        by_category[r["category"]] += r["amount"]
        by_tag[r["tag"] or "(태그 없음)"] += r["amount"]

    if by_period:
        row = _breakdown(ws, row, "주기별", list(by_period.items()))
        row = _breakdown(ws, row, "카테고리별", list(by_category.items()))
        row = _breakdown(ws, row, "태그별", list(by_tag.items()))

    months = defaultdict(int)
    for r in records:
        months[f"{r['dt']:%Y-%m}"] += r["amount"]
    # 추이는 시간순이라야 읽힌다. 금액순으로 늘어놓으면 2026-08, 2026-05, 2026-07 이 된다.
    row = _breakdown(ws, row, "월별 추이", list(months.items()), sort_key=lambda x: x[0])


def build_all(wb, records, settings) -> dict:
    """리포트 시트를 모두 새로 만들고, 터미널 출력용 이번 주 요약을 돌려준다."""
    summary = build_weekly(wb, records, settings)
    build_monthly(wb, records, settings)
    return summary
