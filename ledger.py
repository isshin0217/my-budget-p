"""
ledger.py
파싱된 카드 내역을 엑셀 가계부(.xlsx)에 누적 저장한다.
- 중복 방지: (날짜시간 + 금액 + 가맹점) 이 이미 있으면 건너뜀
- 태그: 사용자가 엑셀에서 직접 적는 열. 다시 실행해도 보존되고,
        같은 가맹점은 다음부터 자동으로 같은 태그가 붙는다(학습).
- 시트: '내역'(원장) + '설정'(월 예산 등)
실행 예시는 run.py 참고.
"""

from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from parser import period_of

HEADERS = ["날짜시간", "금액", "카테고리", "가맹점", "주기", "태그",
           "카드사", "카드뒷자리", "결제유형", "구분"]
WIDTHS = [17, 12, 12, 24, 8, 12, 12, 10, 10, 8]
SHEET = "내역"
SETTINGS = "설정"

HEAD_FILL = PatternFill("solid", fgColor="0A2342")
HEAD_FONT = Font(color="FFFFFF", bold=True)

# 사용자가 손대는 열(태그)은 눈에 띄게
TAG_FILL = PatternFill("solid", fgColor="FFF6E5")

# 설정 기본값. 엑셀 '설정' 시트에서 직접 고치면 된다.
DEFAULT_SETTINGS = [
    ("월 예산", 2000000, "한 달에 쓸 수 있는 총액"),
    ("주 시작 요일", "월", "카드 선결제 주기에 맞추세요 (월~일)"),
    ("월 고정비(수동)", 0, "자동 감지가 잡기 전 임시값. 0이면 자동 감지만 사용"),
]


def _dt_text(v) -> str:
    """날짜시간 칸을 'YYYY-MM-DD HH:MM' 문자열로 되돌린다.

    엑셀에서 그 셀을 한 번 건드리면 openpyxl 이 datetime 으로 읽어온다. 그러면
    str() 결과가 '...19:55:00' 이 되어 우리가 적어둔 '...19:55' 와 안 맞고,
    중복 판정이 통째로 무너져 같은 결제가 두 번 쌓인다. 정렬도 두 종류가
    섞여 엉킨다. 읽을 때도 쓸 때도 이 함수를 거치게 한다.
    """
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return str(v or "").strip()[:16]


def _style_header(ws, widths):
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"


def _new_book() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    _style_header(ws, WIDTHS)
    return wb


def ensure_settings(wb) -> None:
    """'설정' 시트가 없으면 기본값으로 만든다. 있으면 손대지 않는다."""
    if SETTINGS in wb.sheetnames:
        return
    ws = wb.create_sheet(SETTINGS)
    ws.append(["항목", "값", "설명"])
    _style_header(ws, [18, 14, 40])
    for item, value, desc in DEFAULT_SETTINGS:
        ws.append([item, value, desc])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].fill = TAG_FILL


def read_settings(wb) -> dict:
    """'설정' 시트를 dict로 읽는다."""
    out = {item: value for item, value, _ in DEFAULT_SETTINGS}
    if SETTINGS not in wb.sheetnames:
        return out
    for row in wb[SETTINGS].iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and row[1] is not None:
            out[str(row[0]).strip()] = row[1]
    return out


def migrate(ws) -> bool:
    """예전 버전(주기·태그 열 없음)의 시트를 현재 스키마로 옮긴다.

    헤더 이름을 기준으로 값을 다시 배치하므로 열 순서가 달라도 안전하다.
    사용자가 손으로 덧붙인 열은 뒤에 그대로 남긴다 — 우리가 아는 열만 남기고
    자르면, 옆에 적어둔 메모가 실행할 때마다 조용히 사라진다.
    """
    old = [c.value for c in ws[1]]
    if old[:len(HEADERS)] == HEADERS:
        return False

    pos = {name: i for i, name in enumerate(old) if name}
    # 머리글을 하나도 못 알아보면 이건 원장이 아니다. 그대로 진행하면 남의 시트를
    # 헤더만 남기고 갈아엎는다 — 구글시트에서 받은 파일이 딱 그 경우였다.
    if not any(k in pos for k in ("날짜시간", "금액", "가맹점")):
        raise ValueError(
            f"'{ws.title}' 시트의 머리글을 알아볼 수 없습니다: {[c for c in old[:6]]}\n"
            "가계부 원장이 아닌 것 같습니다. 파일을 확인하세요.")

    extra = [name for name in old if name and name not in HEADERS]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rec = {name: row[i] if i < len(row) else "" for name, i in pos.items()}
        rec.setdefault("태그", "")
        rec["주기"] = rec.get("주기") or period_of(rec.get("카테고리") or "미분류")
        # 'or ""' 를 쓰면 금액 0 이 빈칸이 된다. None 일 때만 빈칸으로 둔다.
        rows.append([rec.get(h) if rec.get(h) is not None else "" for h in HEADERS + extra])

    ws.delete_rows(1, ws.max_row)
    ws.append(HEADERS + extra)
    _style_header(ws, WIDTHS)
    for r in rows:
        ws.append(r)
    return True


def _read_existing(ws) -> tuple[set, dict, dict]:
    """이미 있는 행에서 (중복키, 키→태그, 가맹점→태그)를 뽑는다.

    가맹점→태그는 **같은 태그가 2번 이상 일치할 때만** 만든다.
    한 번 본 걸로 단정하면 스타벅스를 데이트로 한 번 태그했다는 이유로
    혼자 마신 커피까지 데이트가 되어버린다. 갈 때마다 성격이 달라지는
    가맹점은 태그가 엇갈리므로 자동으로 붙지 않는다 — 그게 맞는 동작이다.
    """
    keys, by_key = set(), {}
    seen_tags: dict[str, list[str]] = {}
    tag_i = HEADERS.index("태그")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = (_dt_text(row[0]), row[1], str(row[3]))
        keys.add(key)
        tag = str(row[tag_i]).strip() if len(row) > tag_i and row[tag_i] else ""
        if tag:
            by_key[key] = tag
            seen_tags.setdefault(str(row[3]), []).append(tag)

    by_merchant = {m: tags[0] for m, tags in seen_tags.items()
                   if len(tags) >= 2 and len(set(tags)) == 1}
    return keys, by_key, by_merchant


def append_rows(records: list[dict], path: str = "가계부.xlsx") -> tuple[int, Workbook, Path]:
    """레코드들을 엑셀에 추가. (신규 건수, 워크북, 경로)를 반환한다.

    저장은 하지 않는다 — 호출한 쪽에서 리포트까지 만든 뒤 저장한다.
    """
    p = Path(path)
    wb = load_workbook(p) if p.exists() else _new_book()
    # 시트가 없으면 만든다. wb.active 로 넘어가면 안 된다 — '내역' 이 없는 통합문서
    # (구글시트에서 받은 파일이 그렇다)에서 엉뚱한 시트를 원장으로 오해해, migrate 가
    # 그 시트를 머리글만 남기고 통째로 갈아엎었다. 경고도 에러도 없이.
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
    else:
        ws = wb.create_sheet(SHEET)
        ws.append(HEADERS)
        _style_header(ws, WIDTHS)
    ensure_settings(wb)
    migrate(ws)

    seen, tag_by_key, tag_by_merchant = _read_existing(ws)
    added = 0
    for r in records:
        if not r:
            continue
        key = (r["datetime"], r["amount"], r["merchant"])
        if key in seen:
            continue
        ws.append([
            r["datetime"], r["amount"], r["category"], r["merchant"],
            period_of(r["category"]),
            tag_by_merchant.get(r["merchant"], ""),   # 학습된 태그 자동 적용
            r["card"], r["last4"], r["paytype"], r["kind"],
        ])
        seen.add(key)
        added += 1

    _rewrite_sorted(ws, tag_by_key)
    return added, wb, p


def _rewrite_sorted(ws, tag_by_key: dict):
    """날짜순 정렬 + 사용자가 적은 태그 보존 + 서식 재적용."""
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    # 엑셀이 날짜값으로 바꿔둔 셀과 우리가 적은 문자열이 섞이면 정렬이 엉킨다.
    # 문자열로 통일해 두면 정렬도 중복 판정도 한 가지 규칙으로 돈다.
    for r in rows:
        r[0] = _dt_text(r[0])
    rows.sort(key=lambda r: r[0])

    tag_i = HEADERS.index("태그")
    # 사용자가 오른쪽에 덧붙인 열까지 살린다. HEADERS 길이로 자르면 그 열이 사라진다.
    width = max([len(HEADERS)] + [len(r) for r in rows])
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        r = (r + [""] * width)[:width]
        key = (r[0], r[1], str(r[3]))
        if not r[tag_i] and key in tag_by_key:      # 정렬 과정에서 잃지 않도록
            r[tag_i] = tag_by_key[key]
        ws.append(r)
        ws.cell(ws.max_row, 2).number_format = '#,##0'
        ws.cell(ws.max_row, tag_i + 1).fill = TAG_FILL
